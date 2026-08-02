from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Iterable

import streamlit as st
from docx import Document

# MINERVA_DOCUMENTS_IMPORT_BEGIN
try:
    from .config import DOCUMENTOS_DIR, logger
    from .db import executar_query
except (ImportError, ValueError):
    from config import DOCUMENTOS_DIR, logger
    from db import executar_query
# MINERVA_DOCUMENTS_IMPORT_END


# ======================================================================
# MINERVA AI
# Localização de documentos oficiais, downloads (PDF/DOCX) e busca por
# intenção (mapeia a pergunta do usuário para o documento institucional
# correspondente).
# ======================================================================


def caminhos_documento(nome_arquivo: str) -> Iterable[Path]:
    """Retorna caminhos prováveis para localizar os PDFs oficiais.

    Compatível com execução local e com Docker.
    """
    yield DOCUMENTOS_DIR / nome_arquivo
    yield Path("..") / "documentos" / nome_arquivo
    yield Path("/app") / "documentos" / nome_arquivo
    yield Path("/documentos") / nome_arquivo
    yield Path(nome_arquivo)


@st.cache_data(show_spinner=False)
def obter_binario_pdf(nome_arquivo: str) -> bytes | None:
    """Obtém o arquivo PDF original em bytes.

    Cacheado por nome_arquivo: várias mensagens do histórico podem apontar
    para o mesmo documento (ex. 5 perguntas sobre TCC) — sem cache, a leitura
    do arquivo se repetia uma vez por mensagem toda vez que uma sessão
    recarregava o histórico salvo.
    """
    for caminho in caminhos_documento(nome_arquivo):
        if caminho.exists():
            try:
                return caminho.read_bytes()
            except Exception as exc:
                logger.exception("Erro ao ler PDF em %s: %s", caminho, exc)

    return None


def formatar_documento_para_download(texto_bruto: str | None, nome_arquivo: str) -> str:
    """Formata texto bruto extraído do banco para gerar DOCX editável."""
    if not texto_bruto:
        return "Conteúdo indisponível para estruturação."

    if "calendario" in nome_arquivo.lower() or "calendário" in nome_arquivo.lower():
        texto = texto_bruto

        meses = {
            r"J\s*A\s*N\s*E\s*I\s*R\s*O": "\n\n📅 [ JANEIRO ] ───────────────────\n",
            r"F\s*E\s*V\s*E\s*R\s*E\s*I\s*R\s*O": "\n\n📅 [ FEVEREIRO ] ─────────────────\n",
            r"M\s*A\s*R\s*Ç\s*O": "\n\n📅 [ MARÇO ] ────────────────────\n",
            r"A\s*B\s*R\s*I\s*L": "\n\n📅 [ ABRIL ] ────────────────────\n",
            r"M\s*A\s*I\s*O": "\n\n📅 [ MAIO ] ─────────────────────\n",
            r"J\s*U\s*N\s*H\s*O": "\n\n📅 [ JUNHO ] ────────────────────\n",
            r"J\s*U\s*L\s*H\s*O": "\n\n📅 [ JULHO ] ────────────────────\n",
            r"A\s*G\s*O\s*S\s*T\s*O": "\n\n📅 [ AGOSTO ] ───────────────────\n",
            r"S\s*E\s*T\s*E\s*M\s*B\s*R\s*O": "\n\n📅 [ SETEMBRO ] ─────────────────\n",
            r"O\s*U\s*T\s*U\s*B\s*R\s*O": "\n\n📅 [ OUTUBRO ] ──────────────────\n",
            r"N\s*O\s*V\s*E\s*M\s*B\s*R\s*O": "\n\n📅 [ NOVEMBRO ] ─────────────────\n",
            r"D\s*E\s*Z\s*E\s*M\s*B\s*R\s*O": "\n\n📅 [ DEZEMBRO ] ─────────────────\n",
        }

        for regex_mes, substituto in meses.items():
            texto = re.sub(regex_mes, substituto, texto, flags=re.IGNORECASE)

        return texto.strip()

    texto_limpo = re.sub(r"\s+", " ", texto_bruto).strip()
    return re.sub(r"\.\s+([A-ZÁÉÍÓÚÂÊÔÃÕÇ])", r".\n\n\1", texto_limpo)


@st.cache_data(show_spinner=False)
def gerar_docx(texto_estruturado: str) -> bytes:
    """Gera arquivo DOCX real com python-docx.

    Cacheado por texto_estruturado: a geração é uma função pura do texto —
    cachear evita regerar o mesmo DOCX a cada mensagem do histórico que
    aponta para o mesmo documento, e entre reruns/sessões subsequentes.
    """
    doc = Document()
    doc.add_heading("DOCUMENTO OFICIAL EMITIDO PELA MINERVA AI", level=1)

    for paragrafo in texto_estruturado.split("\n"):
        paragrafo = paragrafo.strip()

        if paragrafo:
            doc.add_paragraph(paragrafo)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return buffer.getvalue()


# Mapeia termos do usuário para documentos oficiais.
# Isso é mais confiável do que depender apenas do nome exato do arquivo.
INTENCOES_DOCUMENTOS: dict[str, list[str]] = {
    "CalendarioAcademico2026.pdf": [
        "calendario",
        "calendário",
        "cronograma",
        "prazos",
        "datas",
        "periodo letivo",
        "período letivo",
        "feriado",
        "matricula",
        "matrícula",
        "rematricula",
        "rematrícula",
    ],
    "Resolucao_Estagio.pdf": [
        "estagio",
        "estágio",
        "resolucao estagio",
        "resolução estágio",
        "termo de compromisso",
        "supervisor",
        "relatorio de estagio",
        "relatório de estágio",
        "carga horária estágio",
    ],
    "Matricula_Estagio.pdf": [
        "matricula estagio",
        "matrícula estágio",
        "inscrição estágio",
        "inscricao estagio",
    ],
    "Resolucao_TCC.pdf": [
        "tcc",
        "trabalho de conclusao",
        "trabalho de conclusão",
        "banca",
        "defesa",
        "orientador",
        "monografia",
    ],
    "Matricula_TCC.pdf": [
        "matricula tcc",
        "matrícula tcc",
        "inscrição tcc",
        "inscricao tcc",
    ],
    "Regimento_FCT.pdf": [
        "regimento fct",
        "normas fct",
        "faculdade de computação",
        "faculdade de computacao",
    ],
    "RegimentoGeral.pdf": [
        "regimento geral",
        "normas gerais ufpa",
    ],
    "Regulamento_Graduacao.pdf": [
        "graduacao",
        "graduação",
        "trancamento",
        "disciplina",
        "aproveitamento",
        "segunda chamada",
        "frequência",
        "frequencia",
        "conceito",
        "coeficiente",
    ],
    "Resolucao_Ativ_Complementares.pdf": [
        "atividade complementar",
        "atividades complementares",
        "horas complementares",
    ],
    "Resolucao_Intercambios.pdf": [
        "intercambio",
        "intercâmbio",
        "mobilidade academica",
        "mobilidade acadêmica",
    ],
    "Resolucao_PraticasEnsino.pdf": [
        "pratica de ensino",
        "prática de ensino",
        "praticas de ensino",
        "práticas de ensino",
    ],
}


def normalizar_texto(texto: str) -> str:
    """Normaliza texto para busca por termos."""
    texto = texto.lower().strip()
    texto = re.sub(r"[^\w\sáéíóúâêôãõç]", " ", texto)
    texto = re.sub(r"\s+", " ", texto)

    return texto


def pontuar_documento(query_normalizada: str, termos: list[str]) -> int:
    """Calcula pontuação de compatibilidade entre pergunta e documento."""
    return sum(1 for termo in termos if termo in query_normalizada)


def buscar_contexto_e_arquivo(query: str) -> tuple[str | None, str | None]:
    """Busca documento oficial relacionado à pergunta.

    Fluxo:
    1. Identifica intenção por palavras associadas.
    2. Busca o documento no banco.
    3. Se não encontrar, tenta fallback por nome de arquivo.
    """
    query_lc = normalizar_texto(query)

    melhor_arquivo = None
    maior_pontuacao = 0

    for nome_arquivo, termos in INTENCOES_DOCUMENTOS.items():
        pontuacao = pontuar_documento(query_lc, termos)

        if pontuacao > maior_pontuacao:
            maior_pontuacao = pontuacao
            melhor_arquivo = nome_arquivo

    if melhor_arquivo:
        resultados = executar_query(
            """
            SELECT nome_arquivo, conteudo
            FROM documentos_ufpa
            WHERE nome_arquivo = %s
            LIMIT 1;
            """,
            params=(melhor_arquivo,),
            fetch=True,
        )

        if resultados:
            return resultados[0][0], resultados[0][1]

    palavras = [p for p in re.findall(r"\w+", query_lc) if len(p) > 3]

    if palavras:
        condicoes = " OR ".join(["nome_arquivo ILIKE %s" for _ in palavras])
        params = tuple(f"%{p}%" for p in palavras)

        resultados = executar_query(
            f"""
            SELECT nome_arquivo, conteudo
            FROM documentos_ufpa
            WHERE ({condicoes})
            LIMIT 1;
            """,
            params=params,
            fetch=True,
        )

        if resultados:
            return resultados[0][0], resultados[0][1]

    return None, None


def usuario_pediu_formato(query: str) -> bool:
    """Verifica se o usuário pediu PDF, Word, DOCX ou formato específico."""
    termos = ["extensão", "extensao", "pdf", "word", "doc", "docx", "formato"]
    query_lc = query.lower()

    return any(termo in query_lc for termo in termos)
