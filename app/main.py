
from __future__ import annotations
# MINERVA_PRIORITY_ROUTER_IMPORT_BEGIN
try:
    from .minerva_priority_router import priority_answer
except (ImportError, ValueError):
    from minerva_priority_router import priority_answer
# MINERVA_PRIORITY_ROUTER_IMPORT_END

from minerva_hybrid import responder_minerva
import os
import re
import uuid
import logging
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

import requests
import streamlit as st
import psycopg2
from psycopg2 import pool
from docx import Document

def normalizar_mensagem_historico(msg, role_padrao="assistant"):
    """
    Converte qualquer formato antigo ou novo para:

    {
        "role": "user|assistant|system",
        "content": "texto"
    }

    Compatível com:
    - dict
    - str
    - tuple
    - list
    - objetos com role/content
    """

    role = role_padrao
    content = ""

    if isinstance(msg, dict):

        role = (
            msg.get("role")
            or msg.get("tipo")
            or msg.get("sender")
            or msg.get("autor")
            or role_padrao
        )

        if msg.get("content") is not None:
            content = msg.get("content")

        elif msg.get("conteudo") is not None:
            content = msg.get("conteudo")

        elif msg.get("message") is not None:
            content = msg.get("message")

        elif msg.get("text") is not None:
            content = msg.get("text")

        elif msg.get("resposta") is not None:
            content = msg.get("resposta")

        else:
            content = ""

    elif isinstance(msg, (tuple, list)) and len(msg) >= 2:

        role = msg[0] or role_padrao
        content = msg[1]

    elif isinstance(msg, str):

        role = role_padrao
        content = msg

    elif msg is None:

        role = role_padrao
        content = ""

    else:

        role_obj = getattr(
            msg,
            "role",
            None
        )

        content_obj = getattr(
            msg,
            "content",
            None
        )

        if role_obj is not None:
            role = role_obj

        if content_obj is not None:
            content = content_obj

        else:
            content = str(msg)


    role = str(
        role or role_padrao
    ).strip().lower()


    aliases = {

        "usuario": "user",
        "usuário": "user",
        "human": "user",
        "pergunta": "user",
        "cliente": "user",

        "assistente": "assistant",
        "bot": "assistant",
        "ai": "assistant",
        "ia": "assistant",
        "minerva": "assistant",
        "resposta": "assistant",
    }


    role = aliases.get(
        role,
        role
    )


    if role not in {
        "user",
        "assistant",
        "system"
    }:

        role = (
            role_padrao
            if role_padrao in {
                "user",
                "assistant",
                "system"
            }
            else "assistant"
        )


    return {
        "role": role,
        "content": str(content or "")
    }



def normalizar_lista_mensagens(messages):
    """
    Corrige históricos mistos.

    Exemplo antigo:

    [
        {"role": "user", "content": "Pergunta"},
        "Resposta da Minerva"
    ]

    Resultado:

    [
        {"role": "user", "content": "Pergunta"},
        {"role": "assistant", "content": "Resposta da Minerva"}
    ]
    """

    if messages is None:
        return []


    if isinstance(
        messages,
        (str, dict)
    ):
        messages = [messages]


    try:
        itens = list(messages)

    except Exception:
        itens = [messages]


    resultado = []
    ultimo_role = None


    for item in itens:

        if isinstance(item, dict):

            role_existente = item.get(
                "role"
            )

            if role_existente:
                role_padrao = str(
                    role_existente
                ).lower()

            elif ultimo_role == "user":
                role_padrao = "assistant"

            elif ultimo_role == "assistant":
                role_padrao = "user"

            else:
                role_padrao = "assistant"

        else:

            if ultimo_role == "user":
                role_padrao = "assistant"

            elif ultimo_role == "assistant":
                role_padrao = "user"

            else:
                role_padrao = "assistant"


        msg = normalizar_mensagem_historico(
            item,
            role_padrao=role_padrao
        )


        if msg["content"].strip():

            resultado.append(
                msg
            )

            ultimo_role = msg["role"]


    return resultado



def normalizar_historico_session_state():
    """
    Corrige st.session_state.messages em memória.

    Resolve históricos antigos contendo:
    - strings
    - dicts
    - listas
    - tuplas
    """

    try:

        atuais = st.session_state.get(
            "messages",
            []
        )

        normalizadas = normalizar_lista_mensagens(
            atuais
        )

        st.session_state["messages"] = (
            normalizadas
        )

        return normalizadas

    except Exception:

        return []



def mensagem_role(msg, default="assistant"):
    """
    Retorna role sem assumir que msg seja dict.
    """

    return normalizar_mensagem_historico(
        msg,
        role_padrao=default
    )["role"]



def mensagem_content(msg, default=""):
    """
    Retorna content sem assumir que msg seja dict.
    """

    resultado = normalizar_mensagem_historico(
        msg
    )["content"]

    if resultado is None:
        return default

    return str(resultado)



# Normaliza automaticamente históricos antigos.
normalizar_historico_session_state()











# =============================================================================
# MINERVA AI - FCT/UFPA
# -----------------------------------------------------------------------------
# Aplicação Streamlit para orientação acadêmica baseada em documentos oficiais.
#
# Melhorias aplicadas:
# - Código mais organizado por blocos.
# - Comentários explicando os principais trechos.
# - Credenciais fora do código-fonte.
# - Histórico separado por sessão.
# - Tratamento de erro com rollback.
# - Busca por intenção.
# - Download em PDF oficial e DOCX real.
# - Prompt de segurança para evitar respostas inventadas.
# =============================================================================


# =============================================================================
# LOGGING
# =============================================================================
# O logging registra falhas técnicas no console/container sem exibir stack trace
# para o usuário final.
logging.basicConfig(
    level=os.getenv("LOG_LEVEL", "INFO"),
    format="%(asctime)s [%(levelname)s] %(message)s",
)
logger = logging.getLogger("minerva")


# =============================================================================
# CONFIGURAÇÕES SEGURAS
# =============================================================================
def get_config(nome: str, padrao: Any = None) -> Any:
    """Busca configuração primeiro em st.secrets, depois em variável de ambiente.

    Ordem:
    1. .streamlit/secrets.toml
    2. Variável de ambiente
    3. Valor padrão
    """
    try:
        if nome in st.secrets:
            return st.secrets[nome]
    except Exception:
        pass

    return os.getenv(nome, padrao)


DB_HOST = get_config("DB_HOST", "ufpa_rag_db")
DB_PORT = int(get_config("DB_PORT", "5432"))
DB_NAME = get_config("DB_NAME", "ufpa_rag")
DB_USER = get_config("DB_USER", "admin")
DB_PASSWORD = get_config("DB_PASSWORD", "")

DB_POOL_MIN = int(get_config("DB_POOL_MIN", "1"))
DB_POOL_MAX = int(get_config("DB_POOL_MAX", "10"))

LLM_API_URL = get_config(
    "LLM_API_URL",
    "http://ufpa_rag_llm:8002/v1/chat/completions",
)

DOCUMENTOS_DIR = Path(str(get_config("DOCUMENTOS_DIR", "documentos")))

UFPA_LOGO_URL = get_config(
    "UFPA_LOGO_URL",
    str(Path(__file__).parent / "assets" / "logo-ufpa.png"),
)


# =============================================================================
# CONFIGURAÇÃO DA PÁGINA STREAMLIT
# =============================================================================
# page_icon usa o mesmo logo do cabeçalho (UFPA_LOGO_URL, definido acima) em vez
# do emoji genérico 🏫 — antes a aba do navegador não tinha identidade visual
# com o resto da aplicação.
st.set_page_config(
    page_title="Minerva - FCT/UFPA",
    layout="centered",
    page_icon=UFPA_LOGO_URL,
    initial_sidebar_state="expanded",
)


# =============================================================================
# CSS DA INTERFACE
# =============================================================================
# Bloco único. As iterações anteriores empilhavam 5 blocos <style> que se
# sobrepunham (vários deles nem miravam mais em classes usadas no HTML atual).
# Este bloco estiliza somente o que existe hoje: as divs ".minerva-v2-*", a
# sidebar, os botões e as mensagens do chat.
st.markdown(
    """
    <style>
    :root {
        --minerva-blue: #1E40AF;
        --minerva-blue-soft: #EFF6FF;
        --minerva-border: #E2E8F0;
        --minerva-text: #0F172A;
        --minerva-muted: #64748B;
        --minerva-bg: #F8FAFC;
        --minerva-white: #FFFFFF;
    }

    html, body, .stApp, [data-testid="stAppViewContainer"] {
        background: var(--minerva-bg) !important;
        color: var(--minerva-text) !important;
    }

    [data-testid="stHeader"] {
        background: transparent !important;
    }

    .block-container {
        max-width: 860px !important;
        padding-top: 1.1rem !important;
        padding-bottom: 6rem !important;
    }

    /* Sidebar */
    [data-testid="stSidebar"] {
        background: var(--minerva-white) !important;
        border-right: 1px solid var(--minerva-border);
    }

    [data-testid="stSidebar"] * {
        color: var(--minerva-text) !important;
    }

    /* Bug real: a regra acima ("*" dentro da sidebar) atinge diretamente o
       texto interno do botão "＋ Nova conversa" (que é type="primary", fundo
       azul sólido) e o força para --minerva-text (quase preto) — texto escuro
       em cima de fundo azul escuro, quase ilegível. Precisa de um seletor mais
       específico, incluindo os descendentes do botão, para vencer o "*". */
    [data-testid="stSidebar"] div.stButton > button[kind="primary"],
    [data-testid="stSidebar"] div.stButton > button[kind="primary"] *,
    [data-testid="stSidebar"] div.stButton > button[data-testid="stBaseButton-primary"],
    [data-testid="stSidebar"] div.stButton > button[data-testid="stBaseButton-primary"] * {
        color: var(--minerva-white) !important;
    }

    [data-testid="stSidebar"] .stCaptionContainer {
        color: var(--minerva-muted) !important;
    }

    .minerva-v2-sidebar-brand strong {
        font-size: 1.05rem;
    }

    .minerva-v2-sidebar-brand span {
        display: block;
        color: var(--minerva-muted);
        font-size: 0.85rem;
    }

    .minerva-v2-history-title,
    .minerva-v2-section-title,
    .minerva-v2-suggestions-label {
        font-weight: 700;
        color: var(--minerva-text);
        margin: 10px 0 6px 0;
        font-size: 0.92rem;
    }

    .minerva-v2-active {
        display: flex;
        flex-direction: column;
        gap: 2px;
        padding: 10px 14px;
        border-radius: 12px;
        background: var(--minerva-blue-soft);
        margin: 6px 0 14px 0;
    }

    .minerva-v2-active strong {
        color: var(--minerva-blue);
    }

    .minerva-v2-active span {
        color: var(--minerva-muted);
        font-size: 0.88rem;
    }

    /* Boas-vindas (tela sem mensagens) */
    .minerva-v2-intro {
        border: 1px solid var(--minerva-border);
        border-radius: 18px;
        background: var(--minerva-white);
        padding: 20px 22px;
        margin: 0 0 18px 0;
    }

    .minerva-v2-intro h2 {
        margin: 0 0 6px 0;
        font-size: 1.3rem;
        color: var(--minerva-text);
    }

    .minerva-v2-intro p {
        margin: 0;
        color: var(--minerva-muted);
        line-height: 1.5;
    }

    /* Placeholder do campo de pergunta (st.chat_input) */
    [data-testid="stChatInput"] textarea::placeholder {
        color: var(--minerva-muted);
        opacity: 1;
    }

    /* Botões */
    div.stButton > button {
        border-radius: 12px !important;
        border: 1px solid var(--minerva-border) !important;
        background: var(--minerva-white) !important;
        color: var(--minerva-text) !important;
        font-weight: 500 !important;
        box-shadow: none !important;
        transition: 0.15s ease-in-out !important;
    }

    div.stButton > button:hover {
        border-color: var(--minerva-blue) !important;
        color: var(--minerva-blue) !important;
        background: var(--minerva-blue-soft) !important;
        transform: none !important;
    }

    /* Botão "primary" (categoria ativa, Nova conversa etc.): a regra genérica
       acima forçava fundo/borda brancos em TODO botão com !important, então
       type="primary" nunca aparecia visualmente diferente de "secondary" — a
       categoria selecionada em render_cards_categorias() ficava sem nenhum
       destaque no grid de botões. Seletores redundantes porque o atributo
       "kind" e o data-testid mudaram entre versões do Streamlit. */
    div.stButton > button[kind="primary"],
    div.stButton > button[data-testid="stBaseButton-primary"],
    div.stButton > button[data-testid="baseButton-primary"] {
        border-color: var(--minerva-blue) !important;
        background: var(--minerva-blue) !important;
        color: var(--minerva-white) !important;
    }

    div.stButton > button[kind="primary"]:hover,
    div.stButton > button[data-testid="stBaseButton-primary"]:hover,
    div.stButton > button[data-testid="baseButton-primary"]:hover {
        background: var(--minerva-blue) !important;
        color: var(--minerva-white) !important;
        opacity: 0.92;
    }

    div[data-testid="stExpander"] {
        border: 1px solid var(--minerva-border) !important;
        border-radius: 14px !important;
        box-shadow: none !important;
        background: var(--minerva-white) !important;
    }

    .stDownloadButton button {
        border-radius: 12px !important;
    }

    /* Bolhas do chat: diferencia usuário (destacado) de assistente (neutro) */
    div[data-testid="stChatMessage"] {
        border-radius: 16px !important;
        padding: 10px 16px !important;
        margin-bottom: 10px !important;
    }

    div[data-testid="stChatMessage"]:has([data-testid="stChatMessageAvatarUser"]) {
        background: var(--minerva-blue-soft) !important;
    }

    div[data-testid="stChatMessage"]:has([data-testid="stChatMessageAvatarAssistant"]) {
        background: var(--minerva-white) !important;
        border: 1px solid var(--minerva-border);
    }

    [data-testid="stChatInput"] {
        max-width: 860px;
        margin: auto;
    }

    hr {
        margin: 0.9rem 0 !important;
    }

    /* Imagens nunca devem estourar a coluna/container em que estão
       (o logo do cabeçalho tem width fixo em pixels — sem isso, uma
       coluna mais estreita que o valor fixo faz a imagem vazar). */
    [data-testid="stImage"] img {
        max-width: 100%;
        height: auto;
    }

    /* ------------------------------------------------------------
       RESPONSIVO — MOBILE (celular e tablets estreitos)
       ------------------------------------------------------------ */
    @media (max-width: 768px) {
        .block-container {
            padding-left: 0.85rem !important;
            padding-right: 0.85rem !important;
            padding-top: 0.85rem !important;
        }

        /* Logo do cabeçalho: cap menor pra não competir por espaço
           com o título "Minerva AI" numa tela estreita. */
        [data-testid="stImage"] img {
            max-width: 40px !important;
        }

        h2 {
            font-size: 1.3rem !important;
        }

        .minerva-v2-intro {
            padding: 16px 16px !important;
        }

        .minerva-v2-intro h2 {
            font-size: 1.15rem;
        }

        .minerva-v2-intro p {
            font-size: 0.9rem;
        }

        /* Grades de botões (categorias, perguntas sugeridas) quebram
           em 2 colunas em vez de espremer 3-4 lado a lado. Linhas que
           já tinham só 2 colunas mantêm o mesmo visual de antes. */
        [data-testid="stHorizontalBlock"] {
            flex-wrap: wrap !important;
            row-gap: 0.6rem !important;
        }

        [data-testid="stHorizontalBlock"] > [data-testid="column"] {
            flex: 1 1 45% !important;
            min-width: 45% !important;
            width: 45% !important;
        }

        /* Alvo de toque confortável (mínimo recomendado ~44px). */
        div.stButton > button {
            min-height: 44px !important;
            font-size: 0.85rem !important;
            padding: 0.5rem 0.6rem !important;
            white-space: normal !important;
        }

        [data-testid="stChatInput"] {
            max-width: 100% !important;
        }

        div[data-testid="stChatMessage"] {
            padding: 8px 12px !important;
            font-size: 0.94rem;
        }
    }

    /* Celulares bem estreitos: 1 botão por linha nas grades. */
    @media (max-width: 420px) {
        [data-testid="stHorizontalBlock"] > [data-testid="column"] {
            flex: 1 1 100% !important;
            min-width: 100% !important;
            width: 100% !important;
        }
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# =============================================================================
# RECURSOS EM CACHE
# =============================================================================
@st.cache_resource(show_spinner=False)
def get_http_session() -> requests.Session:
    """Cria sessão HTTP reutilizável para chamadas ao modelo local."""
    session = requests.Session()
    session.headers.update({"Content-Type": "application/json"})
    return session


@st.cache_resource(show_spinner=False)
def get_db_pool() -> pool.SimpleConnectionPool | None:
    """Cria pool de conexões PostgreSQL.

    O pool evita abrir e fechar conexão a cada pergunta.
    """
    try:
        return psycopg2.pool.SimpleConnectionPool(
            minconn=DB_POOL_MIN,
            maxconn=DB_POOL_MAX,
            host=DB_HOST,
            port=DB_PORT,
            database=DB_NAME,
            user=DB_USER,
            password=DB_PASSWORD,
        )
    except Exception as exc:
        logger.exception("Erro ao inicializar pool de banco: %s", exc)
        return None


http_session = get_http_session()
db_pool = get_db_pool()


# =============================================================================
# BANCO DE DADOS
# =============================================================================
def executar_query(
    sql: str,
    params: tuple[Any, ...] | None = None,
    *,
    fetch: bool = False,
    commit: bool = False,
) -> list[tuple[Any, ...]] | None:
    """Executa SQL de forma centralizada.

    Faz rollback em caso de erro e devolve a conexão para o pool.
    """
    if not db_pool:
        logger.error("Pool de banco indisponível.")
        return None

    conn = None
    resultado = None

    try:
        conn = db_pool.getconn()

        with conn.cursor() as cur:
            cur.execute(sql, params)

            if fetch:
                resultado = cur.fetchall()

            if commit:
                conn.commit()

    except Exception as exc:
        if conn:
            conn.rollback()

        logger.exception("Erro ao executar query: %s", exc)
        return None

    finally:
        if conn:
            db_pool.putconn(conn)

    return resultado


def inicializar_banco() -> None:
    """Garante a existência da tabela de histórico.

    A coluna session_id separa conversas de usuários diferentes.
    """
    executar_query(
        """
        CREATE TABLE IF NOT EXISTS historico_minerva (
            id SERIAL PRIMARY KEY,
            session_id TEXT,
            data_hora TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            pergunta TEXT,
            resposta TEXT
        );
        """,
        commit=True,
    )

    executar_query(
        "ALTER TABLE historico_minerva ADD COLUMN IF NOT EXISTS session_id TEXT;",
        commit=True,
    )

    executar_query(
        "UPDATE historico_minerva SET session_id = 'legacy' WHERE session_id IS NULL;",
        commit=True,
    )


def salvar_no_historico(pergunta: str, resposta: str) -> None:
    """Salva pergunta e resposta da sessão atual."""
    executar_query(
        """
        INSERT INTO historico_minerva (session_id, pergunta, resposta)
        VALUES (%s, %s, %s);
        """,
        params=(st.session_state.session_id, pergunta, resposta),
        commit=True,
    )


def limpar_historico_banco() -> bool:
    """Remove apenas o histórico da sessão atual."""
    try:
        executar_query(
            "DELETE FROM historico_minerva WHERE session_id = %s;",
            params=(st.session_state.session_id,),
            commit=True,
        )
        return True

    except Exception as exc:
        logger.exception("Erro ao limpar histórico: %s", exc)
        return False


def carregar_historico_sessao() -> list[tuple[str, str]]:
    """Carrega o histórico persistido da sessão atual."""
    registros = executar_query(
        """
        SELECT pergunta, resposta
        FROM historico_minerva
        WHERE session_id = %s
        ORDER BY id ASC;
        """,
        params=(st.session_state.session_id,),
        fetch=True,
    )

    return registros or []


# =============================================================================
# DOCUMENTOS E DOWNLOADS
# =============================================================================
def caminhos_documento(nome_arquivo: str) -> Iterable[Path]:
    """Retorna caminhos prováveis para localizar os PDFs oficiais.

    Compatível com execução local e com Docker.
    """
    yield DOCUMENTOS_DIR / nome_arquivo
    yield Path("..") / "documentos" / nome_arquivo
    yield Path("/app") / "documentos" / nome_arquivo
    yield Path("/documentos") / nome_arquivo
    yield Path(nome_arquivo)


def obter_binario_pdf(nome_arquivo: str) -> bytes | None:
    """Obtém o arquivo PDF original em bytes."""
    for caminho in caminhos_documento(nome_arquivo):
        if caminho.exists():
            try:
                return caminho.read_bytes()
            except Exception as exc:
                logger.exception("Erro ao ler PDF em %s: %s", caminho, exc)

    return None


def formatar_documento_para_download(texto_bruto: str | None, nome_arquivo: str) -> str:
    """Formata texto bruto extraído do banco para gerar DOCX editável."""
    if not texto_bruto:
        return "Conteúdo indisponível para estruturação."

    if "calendario" in nome_arquivo.lower() or "calendário" in nome_arquivo.lower():
        texto = texto_bruto

        meses = {
            r"J\s*A\s*N\s*E\s*I\s*R\s*O": "\n\n📅 [ JANEIRO ] ───────────────────\n",
            r"F\s*E\s*V\s*E\s*R\s*E\s*I\s*R\s*O": "\n\n📅 [ FEVEREIRO ] ─────────────────\n",
            r"M\s*A\s*R\s*Ç\s*O": "\n\n📅 [ MARÇO ] ────────────────────\n",
            r"A\s*B\s*R\s*I\s*L": "\n\n📅 [ ABRIL ] ────────────────────\n",
            r"M\s*A\s*I\s*O": "\n\n📅 [ MAIO ] ─────────────────────\n",
            r"J\s*U\s*N\s*H\s*O": "\n\n📅 [ JUNHO ] ────────────────────\n",
            r"J\s*U\s*L\s*H\s*O": "\n\n📅 [ JULHO ] ────────────────────\n",
            r"A\s*G\s*O\s*S\s*T\s*O": "\n\n📅 [ AGOSTO ] ───────────────────\n",
            r"S\s*E\s*T\s*E\s*M\s*B\s*R\s*O": "\n\n📅 [ SETEMBRO ] ─────────────────\n",
            r"O\s*U\s*T\s*U\s*B\s*R\s*O": "\n\n📅 [ OUTUBRO ] ──────────────────\n",
            r"N\s*O\s*V\s*E\s*M\s*B\s*R\s*O": "\n\n📅 [ NOVEMBRO ] ─────────────────\n",
            r"D\s*E\s*Z\s*E\s*M\s*B\s*R\s*O": "\n\n📅 [ DEZEMBRO ] ─────────────────\n",
        }

        for regex_mes, substituto in meses.items():
            texto = re.sub(regex_mes, substituto, texto, flags=re.IGNORECASE)

        return texto.strip()

    texto_limpo = re.sub(r"\s+", " ", texto_bruto).strip()
    return re.sub(r"\.\s+([A-ZÁÉÍÓÚÂÊÔÃÕÇ])", r".\n\n\1", texto_limpo)


def gerar_docx(texto_estruturado: str) -> bytes:
    """Gera arquivo DOCX real com python-docx."""
    doc = Document()
    doc.add_heading("DOCUMENTO OFICIAL EMITIDO PELA MINERVA AI", level=1)

    for paragrafo in texto_estruturado.split("\n"):
        paragrafo = paragrafo.strip()

        if paragrafo:
            doc.add_paragraph(paragrafo)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return buffer.getvalue()




# =============================================================================
# BUSCA POR INTENÇÃO
# =============================================================================
# Mapeia termos do usuário para documentos oficiais.
# Isso é mais confiável do que depender apenas do nome exato do arquivo.
INTENCOES_DOCUMENTOS: dict[str, list[str]] = {
    "CalendarioAcademico2026.pdf": [
        "calendario",
        "calendário",
        "cronograma",
        "prazos",
        "datas",
        "periodo letivo",
        "período letivo",
        "feriado",
        "matricula",
        "matrícula",
        "rematricula",
        "rematrícula",
    ],
    "Resolucao_Estagio.pdf": [
        "estagio",
        "estágio",
        "resolucao estagio",
        "resolução estágio",
        "termo de compromisso",
        "supervisor",
        "relatorio de estagio",
        "relatório de estágio",
        "carga horária estágio",
    ],
    "Matricula_Estagio.pdf": [
        "matricula estagio",
        "matrícula estágio",
        "inscrição estágio",
        "inscricao estagio",
    ],
    "Resolucao_TCC.pdf": [
        "tcc",
        "trabalho de conclusao",
        "trabalho de conclusão",
        "banca",
        "defesa",
        "orientador",
        "monografia",
    ],
    "Matricula_TCC.pdf": [
        "matricula tcc",
        "matrícula tcc",
        "inscrição tcc",
        "inscricao tcc",
    ],
    "Regimento_FCT.pdf": [
        "regimento fct",
        "normas fct",
        "faculdade de computação",
        "faculdade de computacao",
    ],
    "RegimentoGeral.pdf": [
        "regimento geral",
        "normas gerais ufpa",
    ],
    "Regulamento_Graduacao.pdf": [
        "graduacao",
        "graduação",
        "trancamento",
        "disciplina",
        "aproveitamento",
        "segunda chamada",
        "frequência",
        "frequencia",
        "conceito",
        "coeficiente",
    ],
    "Resolucao_Ativ_Complementares.pdf": [
        "atividade complementar",
        "atividades complementares",
        "horas complementares",
    ],
    "Resolucao_Intercambios.pdf": [
        "intercambio",
        "intercâmbio",
        "mobilidade academica",
        "mobilidade acadêmica",
    ],
    "Resolucao_PraticasEnsino.pdf": [
        "pratica de ensino",
        "prática de ensino",
        "praticas de ensino",
        "práticas de ensino",
    ],
}


def normalizar_texto(texto: str) -> str:
    """Normaliza texto para busca por termos."""
    texto = texto.lower().strip()
    texto = re.sub(r"[^\w\sáéíóúâêôãõç]", " ", texto)
    texto = re.sub(r"\s+", " ", texto)

    return texto


def pontuar_documento(query_normalizada: str, termos: list[str]) -> int:
    """Calcula pontuação de compatibilidade entre pergunta e documento."""
    return sum(1 for termo in termos if termo in query_normalizada)


def buscar_contexto_e_arquivo(query: str) -> tuple[str | None, str | None]:
    """Busca documento oficial relacionado à pergunta.

    Fluxo:
    1. Identifica intenção por palavras associadas.
    2. Busca o documento no banco.
    3. Se não encontrar, tenta fallback por nome de arquivo.
    """
    query_lc = normalizar_texto(query)

    melhor_arquivo = None
    maior_pontuacao = 0

    for nome_arquivo, termos in INTENCOES_DOCUMENTOS.items():
        pontuacao = pontuar_documento(query_lc, termos)

        if pontuacao > maior_pontuacao:
            maior_pontuacao = pontuacao
            melhor_arquivo = nome_arquivo

    if melhor_arquivo:
        resultados = executar_query(
            """
            SELECT nome_arquivo, conteudo
            FROM documentos_ufpa
            WHERE nome_arquivo = %s
            LIMIT 1;
            """,
            params=(melhor_arquivo,),
            fetch=True,
        )

        if resultados:
            return resultados[0][0], resultados[0][1]

    palavras = [p for p in re.findall(r"\w+", query_lc) if len(p) > 3]

    if palavras:
        condicoes = " OR ".join(["nome_arquivo ILIKE %s" for _ in palavras])
        params = tuple(f"%{p}%" for p in palavras)

        resultados = executar_query(
            f"""
            SELECT nome_arquivo, conteudo
            FROM documentos_ufpa
            WHERE ({condicoes})
            LIMIT 1;
            """,
            params=params,
            fetch=True,
        )

        if resultados:
            return resultados[0][0], resultados[0][1]

    return None, None


def usuario_pediu_formato(query: str) -> bool:
    """Verifica se o usuário pediu PDF, Word, DOCX ou formato específico."""
    termos = ["extensão", "extensao", "pdf", "word", "doc", "docx", "formato"]
    query_lc = query.lower()

    return any(termo in query_lc for termo in termos)


_MINERVA_RENDER_COUNTER = 0

# =============================================================================
# COMPONENTES FINAIS DE UX
# =============================================================================
# Estes componentes substituem listas longas e abas por uma navegação simples:
# 1. Cards de categoria.
# 2. Perguntas rápidas por categoria.
# 3. Histórico clicável.
# 4. Exportação apenas quando houver conversa.
CATEGORIAS_MINERVA: dict[str, dict[str, list[str] | str]] = {
    "Calendário": {
        "icone": "📅",
        "descricao": "Datas, prazos, matrícula e período letivo.",
        "perguntas": [
            "Me mostre o calendário acadêmico de 2026",
            "Quais são os principais prazos do calendário acadêmico?",
            "Quando começa e termina o período letivo?",
            "Quais datas de matrícula constam no calendário?",
        ],
    },
    "Estágio": {
        "icone": "💼",
        "descricao": "Regras, documentos, matrícula e carga horária.",
        "perguntas": [
            "Quais são as regras de estágio?",
            "Como faço a matrícula de estágio?",
            "Quais documentos são necessários para estágio?",
            "Existe carga horária mínima para estágio?",
        ],
    },
    "TCC": {
        "icone": "🎓",
        "descricao": "Regulamento, banca, matrícula e defesa.",
        "perguntas": [
            "Quero saber o regulamento do TCC",
            "Como faço a matrícula do TCC?",
            "Quais regras existem para banca de TCC?",
            "Quais documentos preciso para defender o TCC?",
        ],
    },
    "Cardápio RU": {
        "icone": "🍽️",
        "descricao": "Consulta ao cardápio publicado pela PROAES.",
        "perguntas": [
            "Qual é o cardápio do RU?",
            "Qual é o cardápio de hoje?",
            "O que tem no almoço do RU?",
            "O que tem no jantar do RU?",
        ],
    },
    "Graduação": {
        "icone": "📘",
        "descricao": "Trancamento, disciplina, segunda chamada e aproveitamento.",
        "perguntas": [
            "O que diz o regulamento de graduação?",
            "Como funciona trancamento de disciplina?",
            "Como funciona segunda chamada?",
            "Como funciona aproveitamento de disciplina?",
        ],
    },
    "Normas FCT/UFPA": {
        "icone": "🏛️",
        "descricao": "Regimento da FCT e normas institucionais.",
        "perguntas": [
            "Me mostre o regimento da FCT",
            "Me mostre o regimento geral da UFPA",
            "Quais normas gerais se aplicam aos discentes?",
            "Onde encontro as regras institucionais da faculdade?",
        ],
    },
    "Guias e Manuais": {
        "icone": "📚",
        "descricao": "Tutoriais e formulários oficiais da PROEG (docentes).",
        "perguntas": [
            "Quais guias e manuais a PROEG disponibiliza?",
            "Onde encontro o tutorial do LabINFRA?",
            "Onde encontro o formulário do Avalia Docente?",
            "Onde encontro o tutorial do portal do docente no SIGAA?",
        ],
    },
}


def enviar_pergunta_minerva(pergunta: str) -> None:
    """Envia uma pergunta pronta para o chat."""
    if pergunta:
        st.session_state.messages.append({"role": "user", "content": pergunta})
        st.rerun()


def obter_perguntas_usuario():
    """
    Recupera perguntas sem assumir formato dict.
    """

    mensagens = normalizar_historico_session_state()

    return [
        mensagem_content(msg).strip()

        for msg in mensagens

        if (
            mensagem_role(msg) == "user"
            and mensagem_content(msg).strip()
        )
    ]


def gerar_txt_historico():
    """
    Exporta histórico de forma compatível
    com mensagens antigas e novas.
    """

    mensagens = normalizar_historico_session_state()

    blocos = []

    for msg in mensagens:

        role = mensagem_role(
            msg
        )

        conteudo = mensagem_content(
            msg
        ).strip()

        if not conteudo:
            continue


        if role == "user":

            rotulo = "USUÁRIO"

        elif role == "system":

            rotulo = "SISTEMA"

        else:

            rotulo = "MINERVA"


        blocos.append(
            f"{rotulo}:\n{conteudo}"
        )


    if not blocos:
        return "Histórico vazio."


    return (
        "\n\n".join(blocos)
        + "\n"
    )


def render_historico_clicavel() -> None:
    """Mostra as últimas perguntas como atalhos clicáveis."""
    perguntas = obter_perguntas_usuario()

    if not perguntas:
        return

    recentes = perguntas[-6:]

    for idx, pergunta in enumerate(reversed(recentes)):
        label = pergunta.strip()

        if len(label) > 48:
            label = label[:45] + "..."

        if st.button(
            label,
            key=f"hist_v2_{idx}_{abs(hash(pergunta))}",
            use_container_width=True,
        ):
            enviar_pergunta_minerva(pergunta)



def render_cards_categorias(prefixo: str = "home") -> None:
    """Renderiza categorias compactas e perguntas sugeridas.

    Melhorias:
    - elimina duplicação visual;
    - destaca categoria selecionada;
    - usa keys exclusivas;
    - reduz altura da página.
    """
    global _MINERVA_RENDER_COUNTER

    _MINERVA_RENDER_COUNTER += 1
    render_id = _MINERVA_RENDER_COUNTER

    if "categoria_minerva_ativa" not in st.session_state:
        st.session_state.categoria_minerva_ativa = "Calendário"

    categorias = list(CATEGORIAS_MINERVA.keys())

    st.markdown(
        '<div class="minerva-v2-section-title">'
        'Acessos rápidos'
        '</div>',
        unsafe_allow_html=True,
    )

    # Largura fixa de coluna por linha: antes, a primeira linha (4 categorias)
    # usava st.columns(4) e a segunda (3 categorias) usava st.columns(3) — como
    # cada chamada divide o espaço disponível igualmente, os botões da segunda
    # linha ficavam visivelmente mais largos que os da primeira. Usar sempre o
    # mesmo número de colunas (colunas extras na última linha ficam vazias)
    # mantém a largura dos botões consistente entre as linhas.
    COLUNAS_POR_LINHA = 4

    for inicio_linha in range(0, len(categorias), COLUNAS_POR_LINHA):
        linha = categorias[inicio_linha : inicio_linha + COLUNAS_POR_LINHA]
        cols = st.columns(COLUNAS_POR_LINHA)

        for idx, col in enumerate(cols):
            if idx >= len(linha):
                continue

            categoria = linha[idx]
            ativo = (
                st.session_state.categoria_minerva_ativa
                == categoria
            )

            with col:
                if st.button(
                    categoria,
                    key=f"{prefixo}_{render_id}_cat_{inicio_linha}_{idx}",
                    use_container_width=True,
                    type="primary" if ativo else "secondary",
                ):
                    st.session_state.categoria_minerva_ativa = categoria
                    st.rerun()

    categoria_ativa = (
        st.session_state.categoria_minerva_ativa
    )

    dados = CATEGORIAS_MINERVA[categoria_ativa]
    descricao = dados.get("descricao", "")

    st.markdown(
        f"""
        <div class="minerva-v2-active">
            <strong>{categoria_ativa}</strong>
            <span>{descricao}</span>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="minerva-v2-suggestions-label">'
        'Perguntas sugeridas'
        '</div>',
        unsafe_allow_html=True,
    )

    perguntas = dados.get("perguntas", [])
    cols = st.columns(2)

    for idx, pergunta in enumerate(perguntas):

        with cols[idx % 2]:

            if st.button(
                pergunta,
                key=f"{prefixo}_{render_id}_q_{idx}",
                use_container_width=True,
            ):
                enviar_pergunta_minerva(pergunta)



# =============================================================================
# MODELO LOCAL / FALLBACK
# =============================================================================
# =============================================================================
# ESCOPO INSTITUCIONAL RÍGIDO — FCT/UFPA
# =============================================================================

NOME_UFPA = "Universidade Federal do Pará"
SIGLA_UFPA = "UFPA"
NOME_FCT = "Faculdade de Computação e Telecomunicações"
SIGLA_FCT = "FCT"

PROMPT_INSTITUCIONAL_UFPA = """
Você é a Minerva, assistente virtual acadêmica da
Faculdade de Computação e Telecomunicações (FCT)
da Universidade Federal do Pará (UFPA).

ESCOPO INSTITUCIONAL OBRIGATÓRIO:

1. UFPA significa exclusivamente Universidade Federal do Pará.
2. FCT significa exclusivamente Faculdade de Computação e Telecomunicações.
3. A FCT tratada nesta aplicação pertence ao contexto da UFPA.
4. Toda pergunta acadêmica ambígua deve ser interpretada primeiro no contexto
   da FCT/UFPA.
5. Nunca associe a FCT desta aplicação à UFPB, UFPE, UPE ou a qualquer outra
   universidade.
6. Nunca diga que FCT significa Faculdade de Ciências Técnicas.
7. Nunca diga que a UFPA é Universidade Federal de Pernambuco,
   Universidade Federal da Paraíba ou qualquer outra instituição.
8. Não transfira informações, currículos, duração de cursos, disciplinas,
   regulamentos, calendários, campi ou procedimentos de outras universidades
   para a UFPA.
9. Não invente normas, artigos, resoluções, prazos, datas, requisitos,
   estruturas curriculares, duração de cursos, número de semestres ou
   procedimentos acadêmicos.
10. Quando não houver base institucional suficiente, diga explicitamente:
    "Não localizei informação institucional suficiente na base disponível
    da FCT/UFPA para responder com segurança."
11. Se a pergunta mencionar explicitamente outra universidade, informe
    educadamente que o escopo da Minerva está restrito à FCT/UFPA.
12. Para perguntas como "para que serve o curso", "como funciona o curso",
    "quantos semestres possui" ou "quais disciplinas existem", não invente
    detalhes. Responda apenas com informações institucionais seguras.
13. Antes de concluir cada resposta, confira mentalmente:
    - FCT = Faculdade de Computação e Telecomunicações;
    - UFPA = Universidade Federal do Pará;
    - nenhuma instituição externa foi usada como fonte implícita.

ESTILO:
- Responda em português do Brasil.
- Seja claro, direto e acadêmico.
- Não crie fatos para preencher lacunas.
- Não mencione outras universidades como se fossem a UFPA.
"""


def pergunta_menciona_instituicao_externa(pergunta: str) -> bool:
    """
    Detecta quando o usuário explicitamente direciona a pergunta
    para outra instituição.

    Perguntas sem instituição explícita continuam sendo interpretadas
    no contexto padrão FCT/UFPA.
    """
    texto = pergunta.casefold()

    # Detecta siglas UFxx diferentes de UFPA.
    siglas = re.findall(r"\buf[a-z]{1,4}\b", texto)

    if any(sigla != "ufpa" for sigla in siglas):
        return True

    instituicoes_externas = (
        "universidade federal de pernambuco",
        "universidade federal da paraíba",
        "universidade federal da paraiba",
        "universidade de pernambuco",
        "universidade federal do ceará",
        "universidade federal do ceara",
        "universidade federal do rio grande do norte",
        "universidade federal do rio de janeiro",
        "universidade de são paulo",
        "universidade de sao paulo",
    )

    return any(nome in texto for nome in instituicoes_externas)


def validar_resposta_institucional(resposta: str) -> str:
    """
    Barreira pós-geração.

    Mesmo que o modelo ignore parcialmente o prompt, uma resposta
    contendo associação institucional incompatível não chega ao usuário.
    """
    if not resposta:
        return (
            "Não localizei informação institucional suficiente na base disponível "
            "da FCT/UFPA para responder com segurança."
        )

    texto = resposta.casefold()

    associacoes_invalidas = (
        "universidade federal de pernambuco",
        "universidade federal da paraíba",
        "universidade federal da paraiba",
        "universidade de pernambuco",
        "faculdade de ciências técnicas",
        "faculdade de ciencias tecnicas",
    )

    if any(termo in texto for termo in associacoes_invalidas):
        logger.warning(
            "Resposta bloqueada por associação institucional incompatível: %s",
            resposta,
        )
        return (
            "Não localizei informação institucional suficiente na base disponível "
            "da FCT/UFPA para responder com segurança.\n\n"
            "**Referência institucional correta:**\n"
            "- **FCT:** Faculdade de Computação e Telecomunicações\n"
            "- **UFPA:** Universidade Federal do Pará"
        )

    # Bloqueia siglas UFxx externas produzidas pelo modelo.
    siglas = re.findall(r"\buf[a-z]{1,4}\b", texto)

    if any(sigla != "ufpa" for sigla in siglas):
        logger.warning(
            "Resposta bloqueada por sigla institucional externa: %s",
            resposta,
        )
        return (
            "A Minerva está configurada para responder exclusivamente no contexto "
            "da Faculdade de Computação e Telecomunicações da Universidade Federal "
            "do Pará — FCT/UFPA."
        )

    return resposta


# =============================================================================
# MODELO LOCAL / FALLBACK RESTRITO À FCT/UFPA
# =============================================================================
def consultar_modelo_local(pergunta: str) -> str:
    """
    Consulta a API local compatível com /v1/chat/completions.

    A resposta é submetida a duas barreiras:
    1. prompt institucional rígido;
    2. validação pós-geração.

    Nota: este método recebe o PROMPT JÁ MONTADO (pergunta + contexto do
    RAG + instruções), não a pergunta crua do usuário — por isso a checagem
    de instituição externa não é feita aqui (ver processar_pergunta). O
    próprio texto do PROMPT_INSTITUCIONAL_UFPA cita "UFPB, UFPE, UPE" como
    exemplos do que não confundir, e checar essas palavras aqui gerava
    falso positivo em toda pergunta que chegasse a este ponto.

    Bug real: este método envolvia "pergunta" (que já é o prompt inteiro
    montado por responder_minerva, com regras + contexto RAG + a pergunta
    real) num segundo wrapper ("Considere obrigatoriamente... Pergunta do
    usuário: <prompt inteiro>"). O resultado era um "Pergunta do usuário"
    contendo, ele mesmo, outro bloco "Pergunta:" lá dentro — confuso para
    um modelo pequeno (3B), que podia perder o foco na pergunta real em
    meio a instruções duplicadas. Envia-se "pergunta" direto como conteúdo
    do usuário: ela já está completa.
    """
    payload = {
        "messages": [
            {
                "role": "system",
                "content": PROMPT_INSTITUCIONAL_UFPA,
            },
            {
                "role": "user",
                "content": pergunta,
            },
        ],
        "temperature": 0.0,
        "max_tokens": 450,
        # O Gemma 4 "pensa" antes de responder por padrão (gasta ~400
        # tokens em raciocínio interno numa pergunta de 4 palavras, estourando
        # o max_tokens antes de gerar a resposta de verdade). Modelos sem
        # esse recurso (Qwen, Llama) simplesmente ignoram o campo.
        "chat_template_kwargs": {"enable_thinking": False},
    }

    try:
        response = http_session.post(
            LLM_API_URL,
            json=payload,
            # 12s originais eram menores que o tempo real de resposta do
            # modelo local em CPU (~1-2min com contexto de RAG) — todo
            # request estourava o timeout antes do modelo terminar.
            timeout=240,
        )

        if response.status_code != 200:
            logger.warning(
                "Modelo retornou status %s: %s",
                response.status_code,
                response.text,
            )
            return (
                "Não localizei informação institucional suficiente na base "
                "disponível da FCT/UFPA para responder com segurança."
            )

        data = response.json()

        resposta = (
            data.get("choices", [{}])[0]
            .get("message", {})
            .get("content", "")
            .strip()
        )

        return validar_resposta_institucional(resposta)

    except Exception as exc:
        logger.exception(
            "Erro ao consultar modelo local da Minerva: %s",
            exc,
        )
        return (
            "No momento, não foi possível consultar a base assistida da FCT/UFPA. "
            "Tente novamente mais tarde."
        )

# =============================================================================
# ESTADO DE SESSÃO
# =============================================================================
# Cada navegador/sessão recebe um ID próprio.
# Isso impede que conversas de usuários diferentes se misturem.
if "session_id" not in st.session_state:
    st.session_state.session_id = str(uuid.uuid4())

inicializar_banco()

if "messages" not in st.session_state:
    st.session_state.messages = []

    registros_passados = carregar_historico_sessao()

    for pergunta, resposta in registros_passados:
        st.session_state.messages.append({"role": "user", "content": pergunta})

        nome_arq, texto_raw = buscar_contexto_e_arquivo(pergunta)
        msg_dict = {"role": "assistant", "content": resposta}

        if nome_arq:
            texto_formatado = formatar_documento_para_download(texto_raw, nome_arq)
            msg_dict["download_file"] = nome_arq
            msg_dict["exibir_extensao"] = usuario_pediu_formato(pergunta)
            msg_dict["pdf_bytes"] = obter_binario_pdf(nome_arq)
            msg_dict["docx_bytes"] = gerar_docx(texto_formatado)

        st.session_state.messages.append(normalizar_mensagem_historico(msg_dict, role_padrao="assistant"))


# =============================================================================
# INTERFACE
# =============================================================================
with st.sidebar:
    st.markdown(
        """
        <div class="minerva-v2-sidebar-brand">
            <strong>Minerva AI</strong>
            <span>FCT/UFPA</span>
        </div>
        """,
        unsafe_allow_html=True,
    )

    perguntas_usuario = obter_perguntas_usuario()

    if st.button(
        "＋ Nova conversa",
        key="sidebar_v2_nova_conversa",
        use_container_width=True,
        type="primary",
    ):
        if limpar_historico_banco():
            st.session_state.messages = []
            st.session_state.categoria_minerva_ativa = "Calendário"
            st.rerun()

    # Histórico aparece somente quando existe.
    if perguntas_usuario:

        st.markdown(
            '<div class="minerva-v2-history-title">'
            'Histórico recente'
            '</div>',
            unsafe_allow_html=True,
        )

        render_historico_clicavel()

        st.markdown("---")

        st.download_button(
            "Exportar histórico",
            data=gerar_txt_historico(),
            file_name="historico_minerva.txt",
            mime="text/plain",
            use_container_width=True,
            key="sidebar_v2_export",
        )

    st.markdown("---")

    st.caption(
        "**M.I.N.E.R.V.A.** – Módulo Inteligente de Navegação e "
        "Ensino de Recursos Virtuais Acadêmicos. Assistente "
        "acadêmica baseada em documentos institucionais da FCT/UFPA."
    )




def render_downloads(msg: dict[str, Any], idx: int) -> None:
    """Exibe botões de download do PDF e do DOCX."""
    nome_original = msg["download_file"]
    pede_extensao = msg.get("exibir_extensao", False)

    if pede_extensao:
        caption_arquivo = f"Arquivo oficial: **{nome_original}**"
        label_pdf = "Baixar PDF"
        label_docx = "Baixar Word"
    else:
        nome_formatado = Path(nome_original).stem.replace("_", " ")
        caption_arquivo = f"Documento oficial: **{nome_formatado}**"
        label_pdf = "Baixar PDF"
        label_docx = "Baixar Word"

    st.markdown("---")
    st.caption(caption_arquivo)

    col1, col2 = st.columns(2)

    if msg.get("pdf_bytes"):
        col1.download_button(
            label=label_pdf,
            data=msg["pdf_bytes"],
            file_name=nome_original,
            mime="application/pdf",
            key=f"pdf_btn_{idx}",
        )
    else:
        col1.info("PDF original indisponível.")

    if msg.get("docx_bytes"):
        col2.download_button(
            label=label_docx,
            data=msg["docx_bytes"],
            file_name=nome_original.replace(".pdf", ".docx").replace(".PDF", ".docx"),
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            key=f"docx_btn_{idx}",
        )


def render_chat_history() -> None:
    """Renderiza o histórico completo da conversa."""
    for idx, msg in enumerate(st.session_state.messages):
        avatar = "🏫" if mensagem_role(msg) == "assistant" else "👤"

        with st.chat_message(mensagem_role(msg), avatar=avatar):
            st.markdown(mensagem_content(msg))

            if mensagem_role(msg) == "assistant" and "download_file" in msg:
                render_downloads(msg, idx)



# CABECALHO NATIVO MINERVA V4
# Usa componentes Streamlit para evitar HTML aparecendo como texto.
_header_logo, _header_texto = st.columns([1, 12], gap="small")

with _header_logo:
    st.image(
        UFPA_LOGO_URL,
        width=56,
    )

with _header_texto:
    st.markdown("## Minerva AI")
    st.caption("Orientação acadêmica da FCT/UFPA")

st.divider()
# FIM CABECALHO NATIVO MINERVA V4


# =============================================================================
# ÁREA PRINCIPAL UX V2
# =============================================================================

# -------------------------------------------------------------------------
# 1. Introdução curta
# -------------------------------------------------------------------------
if len(st.session_state.messages) == 0:

    st.markdown(
        """
        <div class="minerva-v2-intro">
            <h2>Olá, sou a Minerva.</h2>
            <p>
                Consulte informações acadêmicas da FCT/UFPA
                ou o cardápio publicado pela PROAES.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )


# -------------------------------------------------------------------------
# 2. Histórico visual da conversa
# -------------------------------------------------------------------------
if st.session_state.messages:
    render_chat_history()


# -------------------------------------------------------------------------
# 3. Categorias
# -------------------------------------------------------------------------
if len(st.session_state.messages) == 0:

    render_cards_categorias(
        prefixo="home_v2"
    )

else:

    with st.expander(
        "Sugestões de perguntas",
        expanded=False,
    ):
        render_cards_categorias(
            prefixo="chat_v2"
        )


# -------------------------------------------------------------------------
# 4. Campo de pergunta — fixo no rodapé e com rolagem automática, nativos
#    do st.chat_input (substitui o antigo st.form + text_input do topo).
# -------------------------------------------------------------------------
pergunta_digitada = st.chat_input(
    "Faça sua pergunta"
)

if pergunta_digitada and pergunta_digitada.strip():

    st.session_state.messages.append(
        {
            "role": "user",
            "content": pergunta_digitada.strip(),
        }
    )

    st.rerun()



# =============================================================================
# PROCESSAMENTO DA PERGUNTA
# =============================================================================
def processar_pergunta(pergunta):
    """
    Processa perguntas da Minerva usando camada institucional FCT/UFPA.
    """
    # MINERVA_PRIORITY_ROUTER_GUARD_BEGIN
    _minerva_priority_result = priority_answer(pergunta)
    if _minerva_priority_result is not None:
        return _minerva_priority_result
    # MINERVA_PRIORITY_ROUTER_GUARD_END

    # Checagem de instituição externa feita aqui, sobre a pergunta CRUA do
    # usuário — e não dentro de consultar_modelo_local, que só recebe o
    # prompt já montado (pergunta + contexto + instruções institucionais).
    if pergunta_menciona_instituicao_externa(pergunta):
        return (
            "Meu escopo institucional está restrito à "
            "**Faculdade de Computação e Telecomunicações (FCT)** da "
            "**Universidade Federal do Pará (UFPA)**. "
            "Não utilizo informações de outras universidades para responder "
            "como se fossem dados da UFPA."
        )

    # Bug real: este trecho procurava por get_connection/conectar_postgres/
    # conectar_banco, nenhuma das quais existe neste módulo (main.py usa
    # db_pool + executar_query). "conn" nunca era preenchido, então
    # buscar_em_documentos_postgres(conn=None, ...) sempre devolvia [] e a
    # Minerva caía direto no fallback genérico — nunca chegava a consultar
    # o banco nem a chamar o modelo local, para nenhuma pergunta fora das
    # rotas fixas do priority_router.
    conn = None

    try:
        if db_pool:
            conn = db_pool.getconn()
    except Exception:
        conn = None

    def llm_func(prompt):
        return consultar_modelo_local(prompt)

    resposta = responder_minerva(
        pergunta=pergunta,
        conn=conn,
        llm_func=llm_func
    )

    if conn:
        try:
            db_pool.putconn(conn)
        except Exception:
            pass

    return resposta


if len(st.session_state.messages) >= 1 and mensagem_role(st.session_state.messages[-1]) == "user":
    ultima_pergunta = mensagem_content(st.session_state.messages[-1])

    with st.chat_message("assistant", avatar="🏫"):
        with st.spinner("Consultando fontes oficiais da FCT/UFPA..."):
            nova_msg = processar_pergunta(ultima_pergunta)

    st.session_state.messages.append(normalizar_mensagem_historico(nova_msg, role_padrao="assistant"))
    st.rerun()
